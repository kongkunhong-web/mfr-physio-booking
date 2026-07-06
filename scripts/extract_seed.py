from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import re

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
ATTACHMENTS = ROOT.parent
MIGRATIONS = ROOT / "migrations"
MIGRATIONS.mkdir(exist_ok=True)

SERVICE_FILES = {
    "ELE": ATTACHMENTS / "ELE.xlsx",
    "GYM": ATTACHMENTS / "GYM.xlsx",
}

SUBTYPES = {
    "ELE": ["ELE-1", "ELE-2", "ELE-3"],
    "GYM": ["GYM-1", "GYM-1/2", "GYM-3", "GYM-3-1"],
}

VALID_LABELS = {
    "ELE": {"ELE": ["ELE-1", "ELE-2", "ELE-3"], "ELE-1": ["ELE-1"]},
    "GYM": {
        "GYM-1": ["GYM-1"],
        "GYM-1/2": ["GYM-1/2"],
        "GYM-3": ["GYM-3"],
        "GYM-3-1": ["GYM-3-1"],
    },
}

WEEKDAYS_BY_SHEET = {
    "ELE-A": [3, 5],
    "ELE-B": [2, 4],
    "GYM-A": [3, 5],
    "GYM-B": [2, 4],
}


def clean_name(value: object) -> str:
    name = str(value).replace("\n", " ").strip()
    return re.sub(r"\s+", " ", name)


def normalize_label(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().upper().replace(" ", "")


def extract_code(name: str) -> str:
    match = re.search(r"\(([^)]+)\)", name)
    if match:
        return match.group(1).strip().split()[0].replace(" ", "")
    return name[:2].upper()


def sql(value: object) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, int):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def is_time_value(value: object) -> str | None:
    if isinstance(value, (int, float)):
        raw = str(int(value))
    elif isinstance(value, str) and value.strip().isdigit():
        raw = value.strip()
    else:
        return None

    if raw == "2026" or len(raw) not in (3, 4):
        return None
    raw = raw.zfill(4)
    return f"{raw[:2]}:{raw[2:]}"


def parse_unavailable_line(text: str):
    text = text.strip()
    date_range = re.search(r"(\d{1,2})/(\d{1,2})(?:-(\d{1,2})/(\d{1,2}))?", text)
    if not date_range:
        return None

    day1, month1, day2, month2 = date_range.groups()
    start = f"2026-{int(month1):02d}-{int(day1):02d}"
    end = f"2026-{int(month2 or month1):02d}-{int(day2 or day1):02d}"
    time_range = re.search(r"(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})", text)
    if time_range:
        h1, m1, h2, m2 = time_range.groups()
        return start, end, f"{int(h1):02d}:{int(m1):02d}", f"{int(h2):02d}:{int(m2):02d}", 0

    return start, end, None, None, 1


def workbook_sheets(service: str):
    return ["ELE-A", "ELE-B"] if service == "ELE" else ["GYM-A", "GYM-B"]


def extract_therapists():
    therapists: list[tuple[str, str, str, str, str]] = []
    seen_names: set[str] = set()
    for service, path in SERVICE_FILES.items():
        workbook = load_workbook(path, data_only=False, read_only=False)
        service_count = 0
        for sheet in workbook_sheets(service):
            worksheet = workbook[sheet]
            for cell in worksheet[2]:
                if cell.column <= 1 or not isinstance(cell.value, str):
                    continue
                name = clean_name(cell.value)
                if not name or name.lower() == "cancel" or name in seen_names:
                    continue
                seen_names.add(name)
                service_count += 1
                gender = "female" if len(therapists) % 2 == 0 else "male"
                therapists.append(
                    (f"{service.lower()}-{service_count:02d}", name, service, extract_code(name), gender)
                )
    return therapists


def extract_time_slots():
    time_slots: dict[str, set[str]] = defaultdict(set)
    for service, path in SERVICE_FILES.items():
        workbook = load_workbook(path, data_only=False, read_only=False)
        for sheet in workbook_sheets(service):
            worksheet = workbook[sheet]
            for row in range(1, worksheet.max_row + 1):
                time = is_time_value(worksheet.cell(row, 1).value)
                if time:
                    time_slots[service].add(time)
    return time_slots


def extract_slot_capacities(time_slots: dict[str, set[str]]):
    slot_caps: dict[tuple[str, str, int, str], int] = defaultdict(int)
    for service, path in SERVICE_FILES.items():
        workbook = load_workbook(path, data_only=False, read_only=False)
        for sheet in workbook_sheets(service):
            worksheet = workbook[sheet]
            therapist_starts = [
                cell.column
                for cell in worksheet[2]
                if cell.column > 1
                and isinstance(cell.value, str)
                and clean_name(cell.value)
                and clean_name(cell.value).lower() != "cancel"
            ]
            time_rows: list[tuple[int, str]] = []
            for row in range(1, worksheet.max_row + 1):
                time = is_time_value(worksheet.cell(row, 1).value)
                if time:
                    time_rows.append((row, time))

            for index, (row, time) in enumerate(time_rows):
                next_row = time_rows[index + 1][0] if index + 1 < len(time_rows) else min(worksheet.max_row + 1, row + 7)
                for start_col in therapist_starts:
                    counts: dict[str, int] = defaultdict(int)
                    category_col = start_col + 1
                    for scan_row in range(row, next_row):
                        label = normalize_label(worksheet.cell(scan_row, category_col).value)
                        if label in VALID_LABELS[service]:
                            for subtype in VALID_LABELS[service][label]:
                                counts[subtype] += 1
                        elif service == "GYM" and label == "GYM1/2":
                            counts["GYM-1/2"] += 1
                    for subtype, count in counts.items():
                        for weekday in WEEKDAYS_BY_SHEET[sheet]:
                            slot_caps[(service, subtype, weekday, time)] = max(
                                slot_caps[(service, subtype, weekday, time)], count
                            )

    for service in ("ELE", "GYM"):
        for subtype in SUBTYPES[service]:
            for weekday in (2, 3, 4, 5):
                for time in sorted(time_slots[service]):
                    slot_caps[(service, subtype, weekday, time)] = max(
                        slot_caps[(service, subtype, weekday, time)], 1
                    )
    return slot_caps


def extract_unavailable(therapists: list[tuple[str, str, str, str, str]]):
    unavailable: list[tuple[str, str, str, str | None, str | None, int, str]] = []
    for service, path in SERVICE_FILES.items():
        workbook = load_workbook(path, data_only=False, read_only=False)
        service_therapists = [therapist for therapist in therapists if therapist[2] == service]
        for sheet in workbook_sheets(service):
            worksheet = workbook[sheet]
            starts = [
                cell.column
                for cell in worksheet[2]
                if cell.column > 1
                and isinstance(cell.value, str)
                and clean_name(cell.value)
                and clean_name(cell.value).lower() != "cancel"
            ]
            for index, start_col in enumerate(starts):
                if index >= len(service_therapists):
                    continue
                raw = worksheet.cell(4, start_col).value
                if not raw:
                    continue
                for line in str(raw).split("\n"):
                    parsed = parse_unavailable_line(line)
                    if parsed:
                        unavailable.append((service_therapists[index][0], *parsed, line.strip()))
    return unavailable


def build_migration():
    therapists = extract_therapists()
    time_slots = extract_time_slots()
    slot_caps = extract_slot_capacities(time_slots)
    unavailable = extract_unavailable(therapists)

    doctors = [
        ("dr-g", "G", "醫生 G", 30),
        ("dr-d", "D", "醫生 D", 30),
        ("dr-a", "A", "醫生 A", 30),
        ("dr-l", "L", "醫生 L", 30),
        ("dr-loi", "LOI", "醫生 LOI", 30),
        ("dr-w", "W", "醫生 W", 30),
        ("dr-n", "N", "醫生 N", 30),
        ("dr-h", "H", "醫生 H", 30),
        ("dr-i", "I", "醫生 I", 30),
    ]
    patients = [
        ("pat-001", "A001", "6661", "28313731", "測試患者一", "dr-loi", "ELE", "ELE-1", "any", 8, "active"),
        ("pat-002", "A002", "8123", "61234567", "測試患者二", "dr-g", "GYM", "GYM-1/2", "female", 6, "active"),
        ("pat-003", "A003", "9345", "62345678", "測試患者三", "dr-d", "ELE", "ELE-2", "male", 10, "draft"),
        ("pat-004", "A004", "7788", "63456789", "測試患者四", "dr-w", "GYM", "GYM-3", "any", 12, "draft"),
        ("pat-005", "A005", "9900", "64567890", "測試患者五", "dr-n", "ELE", "ELE-3", "female", 7, "pending"),
    ]

    lines = [
        "-- Generated from ELE.xlsx/GYM.xlsx structure; no real patient data is imported.",
        "CREATE TABLE IF NOT EXISTS therapists (id TEXT PRIMARY KEY, name TEXT NOT NULL, service_area TEXT NOT NULL, code TEXT NOT NULL, gender TEXT NOT NULL DEFAULT 'unknown', active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);",
        "CREATE TABLE IF NOT EXISTS doctors (id TEXT PRIMARY KEY, code TEXT NOT NULL UNIQUE, name TEXT NOT NULL, quota INTEGER NOT NULL DEFAULT 30, active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);",
        "CREATE TABLE IF NOT EXISTS unavailable_blocks (id TEXT PRIMARY KEY, therapist_id TEXT NOT NULL, start_date TEXT NOT NULL, end_date TEXT NOT NULL, start_time TEXT, end_time TEXT, all_day INTEGER NOT NULL DEFAULT 0, reason TEXT NOT NULL, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);",
        "CREATE TABLE IF NOT EXISTS slot_capacities (id TEXT PRIMARY KEY, therapist_id TEXT, service_area TEXT NOT NULL, subtype TEXT NOT NULL, weekday INTEGER NOT NULL, time TEXT NOT NULL, capacity INTEGER NOT NULL DEFAULT 1, source TEXT NOT NULL DEFAULT 'excel', updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);",
        "CREATE UNIQUE INDEX IF NOT EXISTS uniq_slot_cap_generic ON slot_capacities(COALESCE(therapist_id, ''), service_area, subtype, weekday, time);",
        "CREATE TABLE IF NOT EXISTS activation_batches (id TEXT PRIMARY KEY, label TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'draft', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, activated_at TEXT);",
        "CREATE TABLE IF NOT EXISTS patients (id TEXT PRIMARY KEY, patient_code TEXT NOT NULL, id_number TEXT NOT NULL, phone TEXT NOT NULL, display_name TEXT NOT NULL, doctor_id TEXT NOT NULL, service_area TEXT NOT NULL, subtype TEXT NOT NULL, gender_preference TEXT NOT NULL DEFAULT 'any', session_count INTEGER NOT NULL, status TEXT NOT NULL DEFAULT 'draft', batch_id TEXT, rules_accepted_at TEXT, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, activated_at TEXT, UNIQUE(id_number, phone));",
        "CREATE TABLE IF NOT EXISTS sms_logs (id TEXT PRIMARY KEY, patient_id TEXT NOT NULL, phone TEXT NOT NULL, message TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'sent', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);",
        "CREATE TABLE IF NOT EXISTS bookings (id TEXT PRIMARY KEY, patient_id TEXT NOT NULL UNIQUE, status TEXT NOT NULL DEFAULT 'confirmed', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);",
        "CREATE TABLE IF NOT EXISTS booking_events (id TEXT PRIMARY KEY, booking_id TEXT NOT NULL, patient_id TEXT NOT NULL, therapist_id TEXT NOT NULL, date TEXT NOT NULL, time TEXT NOT NULL, service_area TEXT NOT NULL, subtype TEXT NOT NULL, session_no INTEGER NOT NULL, status TEXT NOT NULL DEFAULT 'confirmed', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, UNIQUE(therapist_id, date, time, patient_id));",
        "CREATE TABLE IF NOT EXISTS help_requests (id TEXT PRIMARY KEY, patient_id TEXT, id_number TEXT, phone TEXT, note TEXT NOT NULL, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);",
        "CREATE INDEX IF NOT EXISTS idx_booking_events_slot ON booking_events(therapist_id, date, time, status);",
        "CREATE INDEX IF NOT EXISTS idx_patients_status ON patients(status);",
        "CREATE INDEX IF NOT EXISTS idx_unavailable_therapist ON unavailable_blocks(therapist_id, start_date, end_date);",
        "",
    ]

    for row in therapists:
        lines.append(f"INSERT OR IGNORE INTO therapists (id,name,service_area,code,gender) VALUES ({','.join(sql(x) for x in row)});")
    for row in doctors:
        lines.append(f"INSERT OR IGNORE INTO doctors (id,code,name,quota) VALUES ({','.join(sql(x) for x in row)});")
    for index, row in enumerate(unavailable[:80], 1):
        lines.append(
            "INSERT OR IGNORE INTO unavailable_blocks "
            "(id,therapist_id,start_date,end_date,start_time,end_time,all_day,reason) VALUES "
            f"({sql(f'unav-{index:03d}')},{','.join(sql(x) for x in row)});"
        )
    for index, ((service, subtype, weekday, time), cap) in enumerate(sorted(slot_caps.items()), 1):
        lines.append(
            "INSERT OR IGNORE INTO slot_capacities "
            "(id,therapist_id,service_area,subtype,weekday,time,capacity,source) VALUES "
            f"({sql(f'cap-{index:04d}')},NULL,{sql(service)},{sql(subtype)},{weekday},{sql(time)},{max(1, min(cap, 3))},'excel');"
        )
    for row in patients:
        lines.append(
            "INSERT OR IGNORE INTO patients "
            "(id,patient_code,id_number,phone,display_name,doctor_id,service_area,subtype,gender_preference,session_count,status,activated_at) VALUES "
            f"({','.join(sql(x) for x in row)}, CASE WHEN {sql(row[-1])}='active' THEN CURRENT_TIMESTAMP ELSE NULL END);"
        )

    lines.extend(
        [
            "INSERT OR IGNORE INTO bookings (id,patient_id,status) VALUES ('book-001','pat-002','confirmed');",
            "INSERT OR IGNORE INTO booking_events (id,booking_id,patient_id,therapist_id,date,time,service_area,subtype,session_no,status) VALUES ('appt-001','book-001','pat-002','gym-03','2026-07-07','09:00','GYM','GYM-1/2',1,'confirmed');",
            "INSERT OR IGNORE INTO booking_events (id,booking_id,patient_id,therapist_id,date,time,service_area,subtype,session_no,status) VALUES ('appt-002','book-001','pat-002','gym-03','2026-07-09','09:00','GYM','GYM-1/2',2,'confirmed');",
            "INSERT OR IGNORE INTO sms_logs (id,patient_id,phone,message,status) VALUES ('sms-001','pat-001','28313731','【物理治療預約】你的線上預約權限已開通，請到 https://mfr.09071247.xyz 選擇時間。','sent');",
        ]
    )

    (MIGRATIONS / "0001_schema.sql").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"therapists={len(therapists)}")
    print(f"slot_capacities={len(slot_caps)}")
    print(f"unavailable_blocks={len(unavailable[:80])}")
    print(MIGRATIONS / "0001_schema.sql")


if __name__ == "__main__":
    build_migration()
